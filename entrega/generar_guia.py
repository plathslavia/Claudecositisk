"""
Llena la Guía 7 (Trabajo con Vistas Master-Detail) conservando el formato original de la UMB.

Uso (desde la raíz del repositorio):
    pip install python-docx
    python3 entrega/generar_guia.py

Toma la plantilla de entrega/plantilla/Guia_7_original.docx y las capturas de
VitalisEPS/capturas/*.png (generadas con `./gradlew capturas` o tomadas del emulador),
y escribe entrega/Guia_7_Master_Detail_Vitalis_EPS.docx.
Si falta una captura, deja en su lugar un aviso para pegarla a mano.
"""

import copy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph

RAIZ = Path(__file__).resolve().parent.parent
PLANTILLA = RAIZ / "entrega" / "plantilla" / "Guia_7_original.docx"
CAPTURAS = RAIZ / "VitalisEPS" / "capturas"
SALIDA = RAIZ / "entrega" / "Guia_7_Master_Detail_Vitalis_EPS.docx"

ASIGNATURA = "DESARROLLO DE APLICACIONES MÓVILES"
INTEGRANTES = ["Andrés García", "Camilo Casallas", "Cristóbal Moncada"]

ARIAL = "Arial"
GOTHIC = "Century Gothic"


# ---------------------------------------------------------------- utilidades XML

def texto(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def buscar(doc, inicio, desde=None):
    """Primer párrafo (en cuerpo o tablas) cuyo texto empieza por `inicio`."""
    encontrado = desde is None
    for p in doc.element.body.iter(qn("w:p")):
        if not encontrado:
            encontrado = p is desde
            continue
        if texto(p).strip().startswith(inicio):
            return p
    raise LookupError(f"No se encontró el párrafo: {inicio!r}")


def rpr(fuente=ARIAL, negrita=False, cursiva=False, color=None, tam=None):
    r = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts")
    f.set(qn("w:ascii"), fuente)
    f.set(qn("w:hAnsi"), fuente)
    f.set(qn("w:cs"), ARIAL)
    r.append(f)
    if negrita:
        r.append(OxmlElement("w:b"))
        r.append(OxmlElement("w:bCs"))
    if cursiva:
        r.append(OxmlElement("w:i"))
        r.append(OxmlElement("w:iCs"))
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        r.append(c)
    if tam:
        for etiqueta in ("w:sz", "w:szCs"):
            s = OxmlElement(etiqueta)
            s.set(qn("w:val"), str(tam))
            r.append(s)
    return r


def corrida(t, **formato):
    r = OxmlElement("w:r")
    r.append(rpr(**formato))
    el = OxmlElement("w:t")
    el.text = t
    el.set(qn("xml:space"), "preserve")
    r.append(el)
    return r


def parrafo(partes, estilo="Prrafodelista", izquierda=0, sangria=None, alinear="both", despues=None, antes=None):
    """`partes`: texto simple o lista de (texto, {formato})."""
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    if estilo:
        s = OxmlElement("w:pStyle")
        s.set(qn("w:val"), estilo)
        ppr.append(s)
    if despues is not None or antes is not None:
        sp = OxmlElement("w:spacing")
        if antes is not None:
            sp.set(qn("w:before"), str(antes))
        if despues is not None:
            sp.set(qn("w:after"), str(despues))
        ppr.append(sp)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(izquierda))
    if sangria:
        ind.set(qn("w:hanging"), str(sangria))
    ppr.append(ind)
    if alinear:
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), alinear)
        ppr.append(jc)
    p.append(ppr)
    if isinstance(partes, str):
        partes = [(partes, {})]
    for t, formato in partes:
        p.append(corrida(t, **formato))
    return p


def insertar_despues(ancla, nuevos):
    for nuevo in nuevos:
        ancla.addnext(nuevo)
        ancla = nuevo
    return ancla


def quitar_vacios_siguientes(p):
    """Elimina los párrafos vacíos que siguen a `p` dentro del mismo contenedor."""
    sig = p.getnext()
    while sig is not None and sig.tag == qn("w:p") and not texto(sig).strip() and not list(sig.iter(qn("w:drawing"))):
        borrar = sig
        sig = sig.getnext()
        borrar.getparent().remove(borrar)


NEG = {"negrita": True}
CUR_GRIS = {"cursiva": True, "color": "595959", "tam": 18}


def respuesta(t, izquierda=720):
    return parrafo([("Respuesta: ", NEG), (t, {})], estilo=None, izquierda=izquierda, despues=120, antes=80)


def vineta(titulo, t, izquierda=720):
    return parrafo([("– ", {}), (titulo, NEG), (t, {})], izquierda=izquierda, sangria=227, despues=60)


# ---------------------------------------------------------------- imágenes

def fila_imagenes(doc, archivos, ancho_pulgadas):
    """Párrafo centrado con una o varias capturas lado a lado (o avisos si faltan)."""
    p = parrafo([], estilo=None, alinear="center", despues=60, antes=120)
    par = Paragraph(p, doc._body)
    faltantes = []
    for i, nombre in enumerate(archivos):
        ruta = CAPTURAS / nombre
        if i > 0:
            par.add_run("   ")
        if ruta.exists():
            par.add_run().add_picture(str(ruta), width=Inches(ancho_pulgadas))
        else:
            faltantes.append(nombre)
    if faltantes:
        for nombre in faltantes:
            p.append(corrida(f"[Pegar aquí la captura {nombre}] ", cursiva=True, color="C00000", tam=18))
    return p


def pie(t):
    return parrafo([(t, CUR_GRIS)], estilo=None, alinear="center", despues=200)


# ---------------------------------------------------------------- contenido

def llenar(doc):
    cuerpo = doc.element.body

    # Encabezado: nombre de la asignatura en el control de contenido.
    titulo = buscar(doc, "GUÍA DE LABORATORIO DE")
    for sdt in cuerpo.iter(qn("w:sdt")):
        alias = sdt.find(f"{qn('w:sdtPr')}/{qn('w:alias')}")
        if alias is not None and alias.get(qn("w:val")) == "ESCRIBIR EL NOMBRE DE LA ASIGNATURA":
            pr = sdt.find(qn("w:sdtPr"))
            marcador = pr.find(qn("w:showingPlcHdr"))
            if marcador is not None:
                pr.remove(marcador)
            contenido = sdt.find(qn("w:sdtContent"))
            for hijo in list(contenido):
                contenido.remove(hijo)
            r = copy.deepcopy(titulo.find(qn("w:r")))
            r.find(qn("w:t")).text = ASIGNATURA
            contenido.append(r)
            integrantes = parrafo(
                [("Integrantes: " + " · ".join(INTEGRANTES), {"fuente": GOTHIC, "tam": 18})],
                estilo=None, alinear="center", despues=0, antes=40,
            )
            sdt.getparent().addnext(integrantes)
            break

    # Tipo de trabajo: Grupal.
    grupal = buscar(doc, "Grupal")
    celda = grupal.getparent()
    siguiente = celda.getnext()
    siguiente.find(qn("w:p")).append(corrida("X", fuente=GOTHIC, negrita=True, tam=18))

    # Introducción: descripción de la solución.
    evidencias = buscar(doc, "Se deben debe incluir como evidencias")
    insertar_despues(evidencias, [
        parrafo([("Solución desarrollada: ", NEG),
                 ("aplicación “Vitalis EPS”, para que los afiliados de una EPS ficticia gestionen sus citas médicas. "
                  "Se construyó en Android Studio con Kotlin Multiplatform y Compose Multiplatform: la interfaz se escribe una sola vez "
                  "y se ejecuta en Android y en escritorio (Cross Platform). Tiene tres vistas: Principal, con la información de la EPS "
                  "y la próxima cita; Maestro, con la lista de citas del afiliado; y Detalle, con toda la información de la cita elegida, "
                  "donde se puede confirmar o cancelar. El código fuente está en la carpeta VitalisEPS del repositorio.", {})],
                estilo=None, izquierda=0, despues=120, antes=160),
    ])

    # Capturas de cada vista.
    principal = buscar(doc, "Pantallazo Vista Principal(Main)")
    quitar_vacios_siguientes(principal)
    insertar_despues(principal, [
        fila_imagenes(doc, ["01_principal.png", "02_principal_empresa.png", "03_principal_contacto.png"], 1.55),
        pie("Figuras 1 a 3. Vista Principal: portada con la próxima cita y el acceso a Mis citas, información de la EPS "
            "(cifras, quiénes somos, misión y visión), servicios y canales de atención."),
    ])

    maestro = buscar(doc, "Pantallazo Vista Maestro")
    quitar_vacios_siguientes(maestro)
    insertar_despues(maestro, [
        fila_imagenes(doc, ["04_maestro.png"], 2.3),
        pie("Figura 4. Vista Maestro (Mis citas): resumen por estado, búsqueda, filtros (Próximas, Por confirmar, "
            "Historial, Todas) y la lista de citas agrupada por mes. Al tocar una cita se abre su detalle."),
    ])

    detalle = buscar(doc, "Pantallazo Vista Detalle")
    quitar_vacios_siguientes(detalle)
    insertar_despues(detalle, [
        fila_imagenes(doc, ["05_detalle.png", "06_detalle_info.png"], 2.3),
        pie("Figuras 5 y 6. Vista Detalle de una cita de Cardiología: profesional, fecha, hora y duración, lugar, motivo, "
            "lista de preparación, autorización y botones para confirmar o cancelar la asistencia."),
        fila_imagenes(doc, ["07_tablet_maestro_detalle.png"], 5.6),
        pie("Figura 7. En pantallas anchas (tablet o escritorio) el Maestro y el Detalle se muestran lado a lado, "
            "como la MasterDetailPage/FlyoutPage de Xamarin."),
    ])

    presentar = buscar(doc, "Presentar la aplicación funcionando al docente")
    insertar_despues(presentar, [
        parrafo("Para la presentación se abre la carpeta VitalisEPS en Android Studio, se espera la sincronización de Gradle, "
                "se elige la configuración composeApp con un emulador o un celular conectado y se pulsa Run. Recorrido de la "
                "demostración: Principal → Ver mis citas → filtrar o buscar → tocar una cita → Detalle → Confirmar asistencia → "
                "botón Atrás del sistema (la cita aparece ahora como Confirmada en la lista).",
                izquierda=720, despues=120, antes=60),
    ])

    # Preguntas orientadoras.
    p1 = buscar(doc, "Está desarrollando una aplicación móvil que requiere del uso de sistema GPS")
    insertar_despues(p1, [respuesta(
        "DependencyService. Es la característica de Xamarin.Forms que permite invocar, desde el código compartido, "
        "funcionalidad nativa de cada plataforma. Se define una interfaz en el proyecto compartido (por ejemplo, IServicioGps con "
        "un método ObtenerUbicacion()), se implementa en cada proyecto nativo (Xamarin.Android con LocationManager y Xamarin.iOS "
        "con CLLocationManager), se registra con el atributo [assembly: Dependency(typeof(ServicioGpsAndroid))] y se obtiene en "
        "tiempo de ejecución con DependencyService.Get<IServicioGps>(). La librería Xamarin.Essentials (Geolocation) ya trae esa "
        "implementación nativa lista. En nuestra aplicación aplicamos el mismo principio con expect/actual de Kotlin Multiplatform: "
        "el código común declara BotonAtras() y BarraDeEstadoClara(), y Android y escritorio aportan cada uno su implementación nativa.")])

    p2 = buscar(doc, "Para el desarrollo de una aplicación móvil que utilizarán personas con discapacidad visual")
    insertar_despues(p2, [respuesta(
        "un servicio de texto a voz (Text-to-Speech) implementado de forma nativa en cada plataforma y consumido con "
        "DependencyService: en Android con la clase Android.Speech.Tts.TextToSpeech, en iOS con AVSpeechSynthesizer y en "
        "Windows con SpeechSynthesizer. Así la lectura usa el motor de voz propio de cada dispositivo, con su idioma y su "
        "configuración de accesibilidad (TalkBack o VoiceOver). Como alternativa, el componente TextToSpeech de "
        "Xamarin.Essentials da ese acceso nativo con una sola llamada: await TextToSpeech.SpeakAsync(texto).")])

    p3 = buscar(doc, "Uno de sus colegas de trabajo tiene problemas con la utilización de los Nuget")
    insertar_despues(p3, [respuesta(
        "la guía no incluye las opciones de esta pregunta, así que explicamos el criterio para elegir. NuGet es el gestor de "
        "paquetes de .NET: un paquete NuGet es una librería que se instala desde el Administrador de paquetes NuGet de Visual "
        "Studio, por ejemplo Xamarin.Forms, Xamarin.Essentials, Newtonsoft.Json, sqlite-net-pcl o Refit. En cambio, "
        "DependencyService, los Custom Renderers y XAML son características del propio framework, y el emulador o el SDK de "
        "Android son herramientas del entorno de desarrollo. Ninguno de ellos es un paquete NuGet, así que la opción que no "
        "corresponde es la que nombre uno de estos elementos (por ejemplo, DependencyService).")])

    # Actividad de trabajo autónomo.
    investigue = buscar(doc, "Investigue 3 aplicaciones informáticas")
    uno = buscar(doc, "1.", desde=investigue)
    dos = buscar(doc, "2.", desde=uno)
    tres = buscar(doc, "3.", desde=dos)
    apps = [
        (uno, "Gmail (Google). ",
         "La bandeja de entrada es la vista maestro: lista de correos con remitente, asunto y un fragmento del mensaje. "
         "Al tocar un correo se abre la vista detalle con el mensaje completo y sus adjuntos. En tablets y en la web ambas vistas "
         "aparecen lado a lado."),
        (dos, "Configuración de Android / Ajustes de iOS. ",
         "La lista de categorías (Wi-Fi, Bluetooth, Pantalla, Batería…) es el maestro y cada opción abre su página de detalle. "
         "En tablets y en el iPad se usan dos paneles: el menú a la izquierda y el detalle a la derecha."),
        (tres, "Software de facturación electrónica (por ejemplo, Siigo o Alegra). ",
         "La lista de facturas emitidas es el maestro; cada factura muestra en el detalle su encabezado (número, fecha, cliente) "
         "y sus líneas de productos. Aquí el patrón maestro-detalle también existe en la base de datos."),
    ]
    for p, nombre, t in apps:
        p.append(corrida(" " + nombre, negrita=True))
        p.append(corrida(t))

    facturacion = buscar(doc, "En una aplicación maestro detalle con gestión de base de datos")
    insertar_despues(facturacion, [
        respuesta("se necesitan como mínimo cuatro tablas. El corazón del proceso es la relación maestro-detalle entre "
                  "Facturas (el maestro, el encabezado) y Movimientos (el detalle, las líneas de la factura): una factura tiene "
                  "muchos movimientos y cada movimiento pertenece a una sola factura. Además se necesitan Clientes y Productos para "
                  "no repetir en cada factura los datos del cliente ni los de cada producto:", izquierda=0),
        vineta("Clientes ", "(id_cliente INTEGER PK autoincremental, nombre TEXT NOT NULL, direccion TEXT, telefono TEXT).",
               izquierda=454),
        vineta("Productos ", "(id_producto INTEGER PK, nombre TEXT, precio_unitario REAL, stock INTEGER).", izquierda=454),
        vineta("Facturas — maestro ", "(id_factura INTEGER PK, id_cliente INTEGER FK, fecha TEXT, total REAL).", izquierda=454),
        vineta("Movimientos — detalle ", "(id_movimiento INTEGER PK, id_factura INTEGER FK, id_producto INTEGER FK, "
               "cantidad INTEGER, precio_unitario_facturado REAL, subtotal REAL).", izquierda=454),
        parrafo("Hay tres llaves foráneas: Facturas.id_cliente → Clientes, Movimientos.id_factura → Facturas y "
                "Movimientos.id_producto → Productos. En Movimientos se guarda el precio_unitario_facturado para que la factura "
                "conserve el precio del día de la venta aunque después cambie el precio en Productos. En una app, la vista Maestro "
                "listaría las facturas y la vista Detalle mostraría los movimientos de la factura elegida.",
                izquierda=0, despues=120),
    ])

    # Actividad de comprobación.
    exponga = buscar(doc, "Exponga los resultados del trabajo autónomo")
    insertar_despues(exponga, [
        parrafo("En la exposición se presentan tres puntos: (1) las tres aplicaciones investigadas usan el patrón maestro-detalle "
                "para ir de una lista general a la información de un elemento y, en pantallas grandes, muestran las dos vistas al "
                "mismo tiempo; (2) el modelo de facturación con sus cuatro tablas y las relaciones uno a muchos entre ellas; y (3) "
                "cómo se aplica el patrón en nuestra app Vitalis EPS: la lista de citas es el maestro y cada cita es el detalle.",
                izquierda=0, despues=120, antes=60),
    ])

    # Procedimiento y metodología.
    procedimiento = buscar(doc, "Procedimiento y Metodología de la práctica")
    diagrama = procedimiento.getparent().getparent().getnext().find(f"{qn('w:tc')}/{qn('w:p')}")
    pasos = [
        ("Análisis. ", "Se definió una EPS ficticia, Vitalis EPS, y el contenido de cada vista: Principal (información de la "
         "EPS y próxima cita), Maestro (lista de citas del afiliado) y Detalle (información completa de una cita)."),
        ("Creación del proyecto. ", "En Android Studio se creó un proyecto Kotlin Multiplatform con Compose Multiplatform y dos "
         "destinos, Android y escritorio (JVM). La interfaz vive en commonMain y se comparte entre plataformas."),
        ("Modelo de datos (data/Datos.kt). ", "El objeto Empresa guarda misión, visión, servicios y canales de atención; la "
         "clase de datos Cita describe cada cita (especialidad, profesional, fecha, hora, sede, estado, preparación y "
         "autorización) y el objeto Agenda contiene nueve citas de distintas especialidades."),
        ("Vista Principal (ui/PantallaPrincipal.kt). ", "Portada con logo, eslogan, una línea de pulso animada y la tarjeta de "
         "la próxima cita con el botón Ver mis citas; luego cifras, quiénes somos, misión y visión, servicios, canales de "
         "atención e integrantes."),
        ("Vista Maestro (ui/PantallaMaestro.kt). ", "Resumen de citas por estado, búsqueda por texto, filtros (Próximas, Por "
         "confirmar, Historial, Todas) y lista con LazyColumn agrupada por mes. Al tocar una cita se navega a su detalle."),
        ("Vista Detalle (ui/PantallaDetalle.kt). ", "Encabezado con el color de la especialidad, profesional, tarjeta con fecha, "
         "hora y duración, lugar o videollamada, motivo, lista de preparación que se puede marcar, datos de autorización y una "
         "barra de acciones que cambia según el estado: confirmar o cancelar la cita. El cambio se refleja también en la lista."),
        ("Navegación (App.kt). ", "Pila de navegación Principal → Maestro → Detalle con transiciones animadas; el botón Atrás del "
         "sistema regresa a la vista anterior. Con 840 dp de ancho o más (tablet o escritorio) el Maestro y el Detalle se "
         "muestran lado a lado."),
        ("Código nativo por plataforma (expect/actual). ", "El botón Atrás y el color de la barra de estado se implementan de "
         "forma nativa en Android, igual que se haría con el DependencyService de Xamarin."),
        ("Pruebas. ", "Se ejecutó la aplicación y se revisaron las tres vistas, la búsqueda, los filtros, la navegación de ida y "
         "vuelta y el diseño de dos paneles; de ahí salen las capturas de la sección de resultados."),
        ("Control de versiones. ", "El código fuente se versionó con Git en GitHub (carpeta VitalisEPS del repositorio)."),
    ]
    nuevos = [parrafo([("Desarrollo de la práctica — aplicación Vitalis EPS", NEG)], estilo=None, alinear="left",
                      despues=80, antes=160)]
    for i, (titulo_paso, t) in enumerate(pasos, start=1):
        nuevos.append(parrafo([(f"{i}. ", NEG), (titulo_paso, NEG), (t, {})], estilo=None, izquierda=340, sangria=340,
                              despues=60))
    insertar_despues(diagrama, nuevos)

    # Criterios de entrega: guion en inglés.
    entrega = buscar(doc, "Exposición en el idioma inglés de la actividad realizada")
    guion = [
        "Hello, we are Andrés García, Camilo Casallas and Cristóbal Moncada. Today we present Vitalis EPS, a cross-platform "
        "mobile app built in Android Studio with Kotlin and Compose Multiplatform, where the members of a health insurance "
        "company manage their medical appointments.",
        "The app follows the master-detail pattern. The main screen introduces the company: our numbers, who we are, our mission "
        "and vision, our services and our support channels. It also shows the next appointment. When the user taps “Ver mis "
        "citas”, the app opens the master view: the list of appointments, grouped by month, that can be searched by doctor or "
        "specialty and filtered by status.",
        "Selecting an appointment opens the detail view. It shows the doctor, the date, time and duration, the place or video "
        "call, the reason for the visit, a checklist to prepare for it and the authorization data. From there the user can "
        "confirm or cancel the appointment, and the change also appears in the list. The system back button always returns to "
        "the previous screen. On wide screens, such as tablets or desktop, the list and the detail are shown side by side.",
        "The user interface is written once and shared between Android and desktop. Platform-specific parts, like the back "
        "button, use Kotlin's expect/actual mechanism, which plays the same role as Xamarin's DependencyService. Thank you.",
    ]
    insertar_despues(entrega, [parrafo([("Guion para la exposición en inglés (English presentation script):", NEG)],
                                       estilo=None, alinear="left", despues=60, antes=160)]
                     + [parrafo([(t, {"cursiva": True})], estilo=None, despues=80) for t in guion])

    # Palabras clave.
    clave = buscar(doc, "Palabras Clave")
    celda_clave = clave.getparent().getparent().getnext().find(f"{qn('w:tc')}/{qn('w:p')}")
    celda_clave.append(corrida(
        "Maestro-detalle (Master-Detail), multiplataforma (Cross Platform), Kotlin Multiplatform, Compose Multiplatform, "
        "Android Studio, navegación, LazyColumn, diseño adaptativo, expect/actual, DependencyService.", tam=18))

    # Referencias consultadas, después de la bibliografía recomendada.
    ultima = buscar(doc, "Napier R, Kumar M")
    referencias = [
        "Microsoft. (s. f.). Xamarin.Forms FlyoutPage. Microsoft Learn. "
        "https://learn.microsoft.com/es-es/xamarin/xamarin-forms/app-fundamentals/navigation/flyoutpage",
        "Microsoft. (s. f.). Xamarin.Forms DependencyService. Microsoft Learn. "
        "https://learn.microsoft.com/es-es/xamarin/xamarin-forms/app-fundamentals/dependency-service/introduction",
        "Microsoft. (s. f.). Xamarin.Essentials: Text-to-Speech. Microsoft Learn. "
        "https://learn.microsoft.com/es-es/xamarin/essentials/text-to-speech",
        "JetBrains. (s. f.). Expected and actual declarations. Kotlin Documentation. "
        "https://kotlinlang.org/docs/multiplatform-expect-actual.html",
        "Android Developers. (s. f.). Build a list-detail layout. "
        "https://developer.android.com/develop/ui/compose/layouts/adaptive/list-detail",
    ]
    insertar_despues(ultima, [parrafo([("Referencias consultadas por el grupo", NEG)], estilo=None, alinear="left",
                                      despues=60, antes=160)]
                     + [parrafo(r, estilo=None, izquierda=454, sangria=454, alinear="left", despues=60) for r in referencias])


def main():
    doc = Document(str(PLANTILLA))
    llenar(doc)
    doc.save(str(SALIDA))
    faltan = [n for n in ["01_principal.png", "02_principal_empresa.png", "03_principal_contacto.png", "04_maestro.png",
                          "05_detalle.png", "06_detalle_info.png", "07_tablet_maestro_detalle.png"]
              if not (CAPTURAS / n).exists()]
    print(f"Guía generada: {SALIDA.relative_to(RAIZ)}")
    if faltan:
        print("Capturas pendientes (quedó un aviso en su lugar): " + ", ".join(faltan))


if __name__ == "__main__":
    main()
